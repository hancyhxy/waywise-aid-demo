"""Build 16:9 slides with undistorted 390x844 phone wireframes.
Requires: python-pptx, playwright, Pillow; playwright install chromium.
Run: python3 build_deck.py (screenshots stay in this deliverable directory).
"""
from pathlib import Path
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageOps, ImageDraw

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'screens'
OUT.mkdir(exist_ok=True)
URL = (ROOT / 'index.html').as_uri()
LIVE = 'https://hancyhxy.github.io/waywise-aid-demo/low-fi/'
# Title, query, click, scroll, purpose, elements, action, limitation
PAGES = [
('Today / Journey overview','','',False,'Understand the current journey before choosing a service.','Origin and destination; service alert; journey-specific priorities; recommendation confidence.','Select a priority or scroll to compare Take, Wait and Switch.','All service conditions and confidence values are fictional.'),
('Today / Compare options','','',True,'Compare the consequences of taking, waiting or switching.','Arrival time; total duration; seating or boarding likelihood; walking; transfers.','Choose an option, or open Why? to inspect its rationale.','Outlined recommendation reflects the demo ranking, not live routing.'),
('Today / Need a seat','','[data-pref="seat"]',True,'Make comfort part of the current decision.','The selected priority has an outlined, underlined state. The recommended route can change.','Toggle Need a seat; compare the updated recommendation.','Journey chips affect demo ranking; preference sliders do not.'),
('Today / Normal service','scenario=normal','',True,'Compare options when the usual service is running normally.','Normal-service scenario; 91% demo confidence; direct and alternative routes.','Compare the next train with waiting or switching.','A scenario variation of Today, not a separate navigation page.'),
('Today / Full bus','scenario=capacity','',True,'Make boarding uncertainty visible before committing.','Capacity scenario; boarding constraints; next bus, following bus and alternative bus.','Compare the risk of waiting with the extra walking of switching.','Boarding likelihood is illustrative rather than a capacity prediction.'),
('Explanation / How it works','','#open-method',False,'Explain the basis of the recommendation.','Current conditions; journey context; confidence and uncertainty.','Close the sheet using the handle, Got it, backdrop or Escape.','Explanatory overlay from the existing demo.'),
('Explanation / Take','','[data-detail="take"]',True,'Explain why taking the next train may suit this journey.','Route-specific rationale alongside context and confidence explanations.','Read the trade-off, then return to the option list.','In this scenario: fastest, but a seat is unlikely.'),
('Explanation / Wait','','[data-detail="wait"]',True,'Explain the potential benefit of waiting.','Queue-clearing rationale and the effect of current context.','Return to the list to compare or choose the waiting option.','The queue-clearing statement is fictional demo content.'),
('Explanation / Switch','','[data-detail="switch"]',True,'Explain the trade-off in changing routes.','Reliability benefit; extra walking; one transfer.','Return to the options and make a decision.','No map or navigation flow exists in the source prototype.'),
('Feedback / Decision saved','','[data-choose="take"]',True,'Acknowledge the selected commute option.','A short confirmation with the chosen action and arrival reference.','Continue reviewing the journey after the message dismisses.','Session-only feedback: no persistent saving or real monitoring.'),
('My week / Office day','screen=trips&day=mon','',False,'View a weekly routine and its outbound and return journeys.','Seven-day selector; office-day journey pair; saved routine templates.','Choose another day; use Today to return to route comparison.','Change journey opens this screen; it does not apply a saved route.'),
('My week / University day','screen=trips&day=tue','',False,'Support different destinations and schedules across the week.','University-day outbound and return journeys; scheduled times.','Select another weekday to inspect its routine.','Wednesday repeats office; Thursday uses later university times.'),
('My week / Remote day','screen=trips&day=fri','',False,'Represent a day without a commute.','Remote-day label; no-commute state; reusable routine templates.','Use + to see the routine-editor placeholder.','Notification pausing is descriptive only in this demo.'),
('My week / Personal trip','screen=trips&day=sat','',False,'Show a flexible weekend journey.','One outbound journey; personal-trip category; departure time.','Switch days or return to Today.','A single-leg routine is retained from the source prototype.'),
('My week / No routine','screen=trips&day=sun','',False,'Show the empty state for an unscheduled day.','No saved Sunday journey; guidance to add one.','Use + to inspect the available editor feedback.','The prototype does not implement a complete create/edit form.'),
('Feedback / Routine editor','screen=trips&day=sun','#add-routine',False,'Indicate the intended next step for adding a routine.','Placeholder feedback describes choosing days and adding journey legs.','Dismisses automatically; Edit week triggers the same feedback.','This is a placeholder, not a working routine editor.'),
('Priorities / Default preferences','screen=prefs','',False,'Express default priorities separately from today’s context.','Arrival time, seat chance and reliability sliders; learning and alert toggles.','Move a slider or switch a toggle to see its visible state.','Controls are visual-only and do not persist or affect route ranking.')
]

def text(slide, s, x,y,w,h,size=18,bold=False):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=0
    p=tf.paragraphs[0]; p.text=s; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor(17,17,17)
    return box

prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
errors=[]
with sync_playwright() as p:
    browser=p.chromium.launch()
    page=browser.new_page(viewport={'width':390,'height':844},device_scale_factor=2)
    page.on('pageerror',lambda e:errors.append(str(e)))
    for i,(title,query,click,scroll,purpose,elements,action,limitation) in enumerate(PAGES,1):
        page.goto(URL+'?export=1&'+query)
        if scroll: page.locator('.screen.active').evaluate('(el)=>el.scrollTop=el.scrollHeight')
        if click: page.locator(click).click()
        page.wait_for_timeout(80)
        image=OUT/f'{i:02d}.png'; page.locator('.phone').screenshot(path=str(image))
        assert Image.open(image).size==(780,1688)
        # Every visible pixel must be neutral: white canvas, black/grey strokes/text only.
        im=Image.open(image).convert('RGB')
        assert all(r==g==b for r,g,b in im.getdata()), f'Colour on screen {i}'
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        text(slide,'WAYWISE / LOW-FIDELITY WIREFRAMES',.55,.25,10,.25,11,True)
        text(slide,title,.55,.69,12,.55,27,True)
        # 390:844 maintained exactly; phone is not stretched to slide dimensions.
        h=5.85; w=h*390/844
        slide.shapes.add_picture(str(image),Inches(1.13),Inches(1.35),width=Inches(w),height=Inches(h))
        for label,body,y in [('PURPOSE',purpose,1.55),('KEY ELEMENTS',elements,2.7),('INTERACTION',action,4.05),('PROTOTYPE SCOPE',limitation,5.38)]:
            text(slide,label,4.6,y,7.8,.28,11,True)
            text(slide,body,4.6,y+.31,7.8,.85,19 if label!='PROTOTYPE SCOPE' else 16)
        text(slide,'390 × 844 portrait viewport · white canvas / outline-only UI',.55,7.22,10,.2,9)
        text(slide,f'{i:02d} / {len(PAGES):02d}',12,7.16,1,.25,11)
        slide.notes_slide.notes_text_frame.text=f'{purpose}\n{elements}\n{action}\nScope: {limitation}\nSource: existing Waywise mid-fidelity demo. Wireframe translation created 2026-09-30, not evidence of an earlier design iteration.\nLive prototype: {LIVE}'
    # Functional regression checks, using a fresh page state.
    page.goto(URL+'?export=1')
    page.locator('[data-pref="seat"]').click()
    assert page.locator('[data-pref="seat"]').get_attribute('aria-pressed')=='true'
    page.locator('[data-detail="wait"]').click()
    assert 'queue' in page.locator('#method-sheet').inner_text()
    page.keyboard.press('Escape')
    page.locator('#open-method').click()
    assert 'Current conditions' in page.locator('#method-sheet .method-item').first.inner_text()
    page.keyboard.press('Escape')
    page.locator('[data-screen="trips"]').click()
    for day in ['mon','tue','wed','thu','fri','sat','sun']:
        page.locator(f'[data-day="{day}"]').click()
        assert page.locator(f'[data-day="{day}"]').get_attribute('aria-selected')=='true'
    page.locator('[data-screen="prefs"]').click()
    page.locator('#time-slider').fill('4'); page.locator('#time-slider').dispatch_event('input')
    assert page.locator('#time-value').inner_text()=='4'
    page.locator('.toggle').first.click()
    assert page.locator('.toggle').first.get_attribute('aria-pressed')=='false'
    assert not errors,errors
    browser.close()
prs.save(ROOT/'Waywise-Low-Fidelity-16x9.pptx')
# Screen contact sheet for visual inspection.
thumbs=[]
for file in sorted(OUT.glob('*.png')):
    im=Image.open(file).convert('RGB'); im.thumbnail((195,422))
    tile=Image.new('RGB',(215,452),'white'); tile.paste(im,(10,20)); ImageDraw.Draw(tile).text((10,3),file.stem,fill='black'); thumbs.append(tile)
contact=Image.new('RGB',(215*6,452*3),'white')
for i,im in enumerate(thumbs):contact.paste(im,((i%6)*215,(i//6)*452))
contact.save(OUT/'contact.jpg')
print(f'Built {len(PAGES)} slides. Verified portrait ratio, monochrome pixels, navigation and controls; no JS errors.')
